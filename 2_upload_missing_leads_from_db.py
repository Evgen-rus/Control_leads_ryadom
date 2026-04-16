import logging
from collections import defaultdict
from logging.handlers import TimedRotatingFileHandler
import os
import random
import re
import sqlite3
import time
from datetime import datetime, timedelta, timezone
from typing import Dict, List

import requests
from dotenv import load_dotenv
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import Workbook

# Период выгрузки (дней) — можно менять
DAYS_LOOKBACK = 3

# Имя файла базы данных в корне проекта
DB_PATH = "lr186.db"

# Для этого скрипта нужен доступ и на чтение, и на запись в Google Sheets
GOOGLE_SHEETS_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

# Время в проекте считаем по МСК (UTC+3)
MSK_TZ = timezone(timedelta(hours=3))

# Повторы при временных ошибках API
MAX_API_RETRIES = 4
API_RETRY_BACKOFF_SECONDS = 1.0

EXPORT_STATUS = "отправил"
EXPORT_FILENAME_TEMPLATE = "data_ryadom_{timestamp}.xlsx"
TELEGRAM_SEND_TIMEOUT_SECONDS = 60

LOG_DIR = "logs"
LOG_FILE = os.path.join(LOG_DIR, "upload_missing_leads_from_db.log")

os.makedirs(LOG_DIR, exist_ok=True)

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

file_handler = TimedRotatingFileHandler(
    LOG_FILE,
    when="midnight",
    backupCount=7,
    encoding="utf-8",
)
file_handler.setFormatter(
    logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
)

console_handler = logging.StreamHandler()
console_handler.setFormatter(
    logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
)

logger.handlers.clear()
logger.addHandler(file_handler)
logger.addHandler(console_handler)


def get_env_required(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise ValueError(f"В переменной окружения {name} пусто или нет значения")
    return value


def create_sheets_service(credentials_file: str):
    if not os.path.exists(credentials_file):
        raise FileNotFoundError(f"Файл credentials не найден: {credentials_file}")
    credentials = service_account.Credentials.from_service_account_file(
        credentials_file, scopes=GOOGLE_SHEETS_SCOPES
    )
    return build("sheets", "v4", credentials=credentials)


def execute_with_retries(request, action_name: str):
    last_error = None
    for attempt in range(1, MAX_API_RETRIES + 1):
        try:
            return request.execute()
        except HttpError as exc:
            status = exc.resp.status if exc.resp else None
            if status in (429, 500, 502, 503, 504):
                last_error = exc
            else:
                raise
        except OSError as exc:
            last_error = exc

        if attempt < MAX_API_RETRIES:
            base_delay = API_RETRY_BACKOFF_SECONDS * (2 ** (attempt - 1))
            sleep_time = random.uniform(0, base_delay)
            logger.warning(
                "Временная ошибка API (%s). Повтор через %.2f сек...",
                action_name,
                sleep_time,
            )
            time.sleep(sleep_time)

    raise RuntimeError(
        f"Не удалось выполнить запрос к API ({action_name}) "
        f"после {MAX_API_RETRIES} попыток: {last_error}"
    )


def normalize_header(header: str) -> str:
    return re.sub(r"\s+", " ", str(header).strip().lower())


def quote_sheet_name(sheet_name: str) -> str:
    escaped_sheet_name = sheet_name.replace("'", "''")
    return f"'{escaped_sheet_name}'"


def column_number_to_letter(column_number: int) -> str:
    letters = []
    current_number = column_number
    while current_number > 0:
        current_number, remainder = divmod(current_number - 1, 26)
        letters.append(chr(65 + remainder))
    return "".join(reversed(letters))


def build_sheet_range(sheet_name: str, cell_range: str) -> str:
    return f"{quote_sheet_name(sheet_name)}!{cell_range}"


def get_pending_leads(conn: sqlite3.Connection, cutoff_dt: datetime) -> List[sqlite3.Row]:
    conn.row_factory = sqlite3.Row

    # event_dt хранится в SQLite строкой в формате YYYY-MM-DD HH:MM:SS,
    # поэтому по нему можно надежно фильтровать и сортировать как по тексту.
    cursor = conn.execute(
        """
        SELECT
            source_id,
            event_dt,
            phone,
            sheet_name,
            sheet_row
        FROM leads
        WHERE event_dt >= ?
          AND TRIM(COALESCE(skorozvon_info, '')) = ''
        ORDER BY event_dt ASC, sheet_row ASC
        """,
        (cutoff_dt.strftime("%Y-%m-%d %H:%M:%S"),),
    )
    return cursor.fetchall()


def validate_pending_leads(leads: List[sqlite3.Row]) -> None:
    for lead in leads:
        if not lead["sheet_name"]:
            raise ValueError(
                f"Для source_id={lead['source_id']} не заполнено поле sheet_name"
            )
        if not lead["sheet_row"]:
            raise ValueError(
                f"Для source_id={lead['source_id']} не заполнено поле sheet_row"
            )


def get_sheet_headers(service, spreadsheet_id: str, sheet_name: str) -> List[str]:
    request = (
        service.spreadsheets()
        .values()
        .get(
            spreadsheetId=spreadsheet_id,
            range=build_sheet_range(sheet_name, "1:1"),
        )
    )
    result = execute_with_retries(request, f"get_sheet_headers:{sheet_name}")
    values = result.get("values", [])
    if not values:
        raise ValueError(f"Вкладка {sheet_name} пустая, не удалось прочитать заголовки")
    return values[0]


def get_status_column_number(headers: List[str]) -> int:
    for idx, header in enumerate(headers, start=1):
        if normalize_header(header) == "статус отправки в скорозвон":
            return idx
    raise ValueError("Не найден столбец 'Статус отправки в скорозвон'")


def build_sheet_updates(
    service,
    spreadsheet_id: str,
    leads: List[sqlite3.Row],
) -> List[Dict[str, List[List[str]]]]:
    updates: List[Dict[str, List[List[str]]]] = []
    leads_by_sheet: Dict[str, List[sqlite3.Row]] = defaultdict(list)

    for lead in leads:
        leads_by_sheet[str(lead["sheet_name"])].append(lead)

    for sheet_name, sheet_leads in leads_by_sheet.items():
        headers = get_sheet_headers(service, spreadsheet_id, sheet_name)
        status_column_number = get_status_column_number(headers)
        status_column_letter = column_number_to_letter(status_column_number)

        for lead in sheet_leads:
            range_name = build_sheet_range(
                sheet_name, f"{status_column_letter}{lead['sheet_row']}"
            )
            updates.append(
                {
                    "range": range_name,
                    "values": [[EXPORT_STATUS]],
                }
            )

    return updates


def create_export_file(leads: List[sqlite3.Row]) -> str:
    timestamp = datetime.now(MSK_TZ).strftime("%Y-%m-%d_%H-%M-%S")
    filename = EXPORT_FILENAME_TEMPLATE.format(timestamp=timestamp)
    file_path = os.path.abspath(filename)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Данные"

    worksheet.append(["Дата", "Номера"])
    for lead in leads:
        worksheet.append(
            [str(lead["event_dt"]), normalize_phone_for_export(lead["phone"])]
        )

    for cell in worksheet["B"][1:]:
        cell.number_format = "0"

    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 18

    workbook.save(file_path)
    return file_path


def normalize_phone_for_export(phone: object) -> int:
    phone_text = str(phone).strip()
    normalized_phone = phone_text.lstrip("'").strip()
    return int(normalized_phone)


def send_file_to_telegram(
    bot_token: str,
    chat_id: str,
    file_path: str,
    leads_count: int,
) -> None:
    url = f"https://api.telegram.org/bot{bot_token}/sendDocument"
    caption = f"@limeunicorn Выгрузка данных для скорозвона: {leads_count}"

    with open(file_path, "rb") as document_file:
        response = requests.post(
            url,
            data={
                "chat_id": chat_id,
                "caption": caption,
            },
            files={
                "document": (
                    os.path.basename(file_path),
                    document_file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            timeout=TELEGRAM_SEND_TIMEOUT_SECONDS,
        )

    response.raise_for_status()
    payload = response.json()
    if not payload.get("ok"):
        raise RuntimeError(
            f"Telegram API вернул ошибку при отправке файла: {payload}"
        )


def update_google_sheet_statuses(
    service,
    spreadsheet_id: str,
    updates: List[Dict[str, List[List[str]]]],
) -> None:
    if not updates:
        return

    request = service.spreadsheets().values().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={
            "valueInputOption": "RAW",
            "data": updates,
        },
    )
    execute_with_retries(request, "batch_update_sheet_statuses")


def mark_leads_as_sent(conn: sqlite3.Connection, source_ids: List[str]) -> None:
    if not source_ids:
        return

    rows_to_update = [(EXPORT_STATUS, source_id) for source_id in source_ids]
    conn.executemany(
        """
        UPDATE leads
        SET skorozvon_info = ?
        WHERE source_id = ?
        """,
        rows_to_update,
    )
    conn.commit()


def main() -> None:
    load_dotenv()

    bot_token = get_env_required("TELEGRAM_BOT_TOKEN")
    chat_id = get_env_required("TELEGRAM_LB_CHAT_ID")
    spreadsheet_id = get_env_required("GOOGLE_SHEET_ID")
    credentials_file = get_env_required("GOOGLE_CREDENTIALS_FILE")

    if not re.fullmatch(r"[A-Za-z0-9-_]{20,}", spreadsheet_id):
        raise ValueError(
            "GOOGLE_SHEET_ID должен быть только ID таблицы без URL"
        )

    cutoff_dt = datetime.now(MSK_TZ) - timedelta(days=DAYS_LOOKBACK)

    conn = sqlite3.connect(DB_PATH)
    try:
        leads = get_pending_leads(conn, cutoff_dt)
        if not leads:
            logger.info("Нет новых лидов для выгрузки.")
            return

        validate_pending_leads(leads)

        service = create_sheets_service(credentials_file)
        sheet_updates = build_sheet_updates(service, spreadsheet_id, leads)

        export_file_path = create_export_file(leads)
        logger.info("Файл выгрузки создан: %s", export_file_path)

        send_file_to_telegram(
            bot_token=bot_token,
            chat_id=chat_id,
            file_path=export_file_path,
            leads_count=len(leads),
        )
        logger.info("Файл успешно отправлен в Telegram: %s", os.path.basename(export_file_path))

        update_google_sheet_statuses(service, spreadsheet_id, sheet_updates)
        logger.info("Статусы в Google Sheets обновлены: %s", len(sheet_updates))

        source_ids = [str(lead["source_id"]) for lead in leads]
        mark_leads_as_sent(conn, source_ids)
        logger.info("Статусы в БД обновлены: %s", len(source_ids))

        os.remove(export_file_path)
        logger.info(
            "Файл выгрузки удален после успешной обработки: %s",
            os.path.basename(export_file_path),
        )
    finally:
        conn.close()


if __name__ == "__main__":
    main()
